#!/usr/bin/env python3
"""Mark a prospect row as sent in Appinara_Prospekti_a_Outreach.xlsx.

Usage: python3 mark_sent.py <row> <email> [status]
Writes K=Status, L=Posledny kontakt=today, N=Povodne odoslane=today (if empty).
Verifies the email in column E matches before writing. Saves immediately.
"""
import sys
import datetime
import openpyxl

PATH = "/sessions/beautiful-gracious-brown/mnt/biznis/Appinara_Prospekti_a_Outreach.xlsx"


def main():
    row = int(sys.argv[1])
    email = sys.argv[2].strip().lower()
    status = sys.argv[3] if len(sys.argv) > 3 else "Odoslané"
    today = datetime.date.today().strftime("%d.%m.%Y")

    wb = openpyxl.load_workbook(PATH)
    ws = wb["Prospekti"]

    actual = str(ws.cell(row, 5).value or "").strip().lower()
    if actual != email:
        print(f"ABORT: row {row} email mismatch: sheet={actual!r} expected={email!r}")
        sys.exit(1)

    sender = str(ws.cell(row, 15).value or "").strip()
    if sender != "appinara.sk":
        print(f"ABORT: row {row} sender={sender!r} is not appinara.sk")
        sys.exit(1)

    ws.cell(row, 11).value = status
    ws.cell(row, 12).value = today
    if not str(ws.cell(row, 14).value or "").strip():
        ws.cell(row, 14).value = today

    wb.save(PATH)
    print(
        f"OK row={row} {ws.cell(row, 3).value} | {email} | "
        f"K={ws.cell(row, 11).value} L={ws.cell(row, 12).value} N={ws.cell(row, 14).value}"
    )


if __name__ == "__main__":
    main()
