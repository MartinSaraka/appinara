import openpyxl, sys, datetime
P="Appinara_Prospekti_a_Outreach.xlsx"
TODAY="25.09.2026"
def main():
    row=int(sys.argv[1]); status=sys.argv[2]; note=sys.argv[3] if len(sys.argv)>3 else ""
    wb=openpyxl.load_workbook(P); ws=wb["Prospekti"]
    firma=ws.cell(row,3).value; email=ws.cell(row,5).value
    if status=="Odoslané":
        ws.cell(row,11).value="Odoslané"
        ws.cell(row,12).value=TODAY
        if not (ws.cell(row,14).value and str(ws.cell(row,14).value).strip()):
            ws.cell(row,14).value=TODAY
    elif status:
        ws.cell(row,11).value=status
    if note:
        old=ws.cell(row,13).value or ""
        ws.cell(row,13).value=(old+" | " if old.strip() else "")+note
    wb.save(P)
    print(f"OK r{row} | {firma} | {email} | K={ws.cell(row,11).value} | L={ws.cell(row,12).value} | N={ws.cell(row,14).value}")
main()
