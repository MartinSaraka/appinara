#!/usr/bin/env python3
"""Append a note to column M (Poznamky) for a prospect row, verifying the email first.

Usage: python3 add_note.py <row> <email> "<note text>"
"""
import sys
import openpyxl

PATH = "/sessions/beautiful-gracious-brown/mnt/biznis/Appinara_Prospekti_a_Outreach.xlsx"


def main():
    row = int(sys.argv[1])
    email = sys.argv[2].strip().lower()
    note = sys.argv[3]

    wb = openpyxl.load_workbook(PATH)
    ws = wb["Prospekti"]

    actual = str(ws.cell(row, 5).value or "").strip().lower()
    if actual != email:
        print(f"ABORT: row {row} email mismatch: sheet={actual!r} expected={email!r}")
        sys.exit(1)

    existing = str(ws.cell(row, 13).value or "").strip()
    ws.cell(row, 13).value = (existing + " | " + note).strip(" |") if existing else note
    wb.save(PATH)
    print(f"OK row={row} {ws.cell(row, 3).value} | M={ws.cell(row, 13).value}")


if __name__ == "__main__":
    main()
